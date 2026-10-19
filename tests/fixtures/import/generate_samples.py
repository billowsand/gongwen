#!/usr/bin/env python3
"""生成导入结构清洗用的真实样例。

跑法（仓库根目录）：

    python3 tests/fixtures/import/generate_samples.py
    typst compile tests/fixtures/import/08-fangxun-notice.typ tests/fixtures/import/08-fangxun-notice.pdf

生成物（docx / txt / pdf）与同名 .expected.md 一起提交；改规则后重跑
cargo test --locked realistic_import_samples 对照。
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent


def paragraph(doc, text="", *, bold_lead=None, style=None):
    """加一段；bold_lead 是段首加粗的那截文字。"""
    p = doc.add_paragraph(style=style)
    if bold_lead is not None and text.startswith(bold_lead):
        run = p.add_run(bold_lead)
        run.bold = True
        p.add_run(text[len(bold_lead) :])
    else:
        p.add_run(text)
    return p


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value
    return table


def fangxun_notice():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "仿宋"
    style.font.size = Pt(12)

    # 版头（编辑稿常把版头放在正文里，导入时应被剥掉）。
    paragraph(doc, "秘密★1年")
    paragraph(doc, "×政发〔2026〕12号")

    # 标题在 Word 里被一个软回车折成两行。
    title = doc.add_paragraph()
    title.add_run("关于做好2026年防汛")
    title.add_run().add_break()
    title.add_run("工作的通知")

    paragraph(doc, "各区县人民政府，市政府各部门：")
    paragraph(
        doc,
        "为切实做好2026年防汛工作，保障人民群众生命财产安全，"
        "根据市委、市政府工作部署，现将有关事项通知如下。",
    )
    paragraph(doc, "一、总体要求", bold_lead="一、总体要求")
    paragraph(
        doc,
        "坚持人民至上、生命至上，立足防大汛、抢大险、救大灾，"
        "全面排查整治风险隐患，确保安全度汛。",
    )
    paragraph(
        doc,
        "（一）加强组织领导。各单位要成立防汛工作领导小组，"
        "明确责任分工，主要负责同志亲自抓、负总责。",
        bold_lead="（一）加强组织领导。",
    )
    paragraph(doc, "（二）压实工作责任")
    paragraph(doc, "1.排查隐患")
    paragraph(doc, "对河道、水库、地质灾害隐患点开展拉网式排查，建立问题台账。")
    paragraph(doc, "2.整改落实")
    paragraph(doc, "对排查发现的问题限期整改，重大隐患实行挂牌督办。")
    paragraph(doc, "二、保障措施")
    paragraph(doc, "1.加强值班值守；")
    paragraph(doc, "2.强化物资储备；")
    paragraph(doc, "3.严格信息报送。")
    paragraph(doc, "各单位要于2026年5月20日前将工作情况报送市防汛抗旱指挥部办公室。")
    paragraph(doc, "附件：1.防汛物资储备清单 2.隐患排查责任分工表")

    # 附件一：普通表格。
    paragraph(doc, "附件1")
    paragraph(doc, "防汛物资储备清单")
    add_table(
        doc,
        [
            ["序号", "物资名称", "单位", "数量"],
            ["1", "编织袋", "万条", "20"],
            ["2", "冲锋舟", "艘", "15"],
            ["3", "救生衣", "件", "3000"],
        ],
    )
    paragraph(doc, "附件2")
    paragraph(doc, "隐患排查责任分工表")
    table = add_table(
        doc,
        [
            ["责任单位", "排查范围", "完成时限"],
            ["市应急管理局", "河道、水库", "2026年5月底"],
            ["市住房城乡建设局", "在建工程", "2026年5月底"],
        ],
    )
    # 合并最后一行的前两格，验证转换器只留空格子时我们的输出。
    table.cell(2, 0).merge(table.cell(2, 1))

    # 落款、成文日期、版记。
    paragraph(doc, "××市人民政府办公室")
    paragraph(doc, "2026年5月10日")
    paragraph(doc, "抄送：市委办公室，市人大常委会办公室。")
    paragraph(doc, "××市人民政府办公室 2026年5月10日印发")

    doc.save(ROOT / "01-fangxun-notice.docx")


def approval():
    doc = Document()
    paragraph(doc, "关于申请增加防汛物资储备经费的请示")
    paragraph(doc, "市政府：")
    paragraph(
        doc,
        "当前已进入主汛期，现有防汛物资储备难以满足抢险需要。"
        "一是部分物资储备不足，缺口较大；二是部分物资老化严重，需要更新；"
        "三是储备库点布局不合理，调运时间较长。",
    )
    paragraph(doc, "一、基本情况")
    paragraph(
        doc,
        "全市现有防汛物资储备库12处，储备编织袋20万条、冲锋舟15艘、救生衣3000件。",
    )
    paragraph(
        doc,
        "近期排查发现，部分区县物资储备不足，个别物资已超过使用年限，"
        "难以满足防大汛、抢大险需要。",
    )
    paragraph(
        doc,
        "拟申请经费500万元，主要用于补充编织袋、冲锋舟、救生衣等物资，"
        "并对现有储备库进行维修改造。",
    )
    paragraph(doc, "妥否，请指示。")
    paragraph(doc, "××市应急管理局")
    paragraph(doc, "2026年5月8日")

    doc.save(ROOT / "02-approval.docx")


def regulation():
    doc = Document()
    paragraph(doc, "××市防汛物资管理办法")
    paragraph(doc, "第一章 总则")
    paragraph(doc, "第一条 为加强防汛物资管理，保障防汛抢险需要，根据有关法律法规，制定本办法。")
    paragraph(doc, "第二条 本办法适用于本市行政区域内防汛物资的储备、调运、使用和管理。")
    paragraph(doc, "第二章 储备管理")
    paragraph(doc, "第三条 市应急管理部门负责统筹全市防汛物资储备工作，各区县应急管理部门负责本辖区防汛物资储备管理。")
    paragraph(doc, "第四条 储备单位应当建立台账，定期盘点，确保账实相符。")
    paragraph(doc, "第三章 调运使用")
    paragraph(doc, "第五条 防汛物资调用实行分级负责、就近调运原则，任何单位和个人不得擅自挪用。")
    paragraph(doc, "第六条 本办法自2026年6月1日起施行。")

    doc.save(ROOT / "03-regulation.docx")


def styled_plan():
    doc = Document()
    doc.add_heading("关于2026年防汛工作的实施方案", level=1)
    doc.add_heading("一、工作目标", level=1)
    doc.add_heading("（一）总体目标", level=2)
    paragraph(
        doc,
        "到2027年基本建成覆盖全市的防汛减灾体系，重点领域和关键环节取得实质性进展。",
    )
    doc.add_heading("（二）具体指标", level=2)
    paragraph(
        doc,
        "病险水库除险加固率达到100%，城市易涝点整治率达到95%以上。",
    )
    doc.add_heading("二、重点任务", level=1)
    doc.add_heading("（一）排查整治", level=2)
    paragraph(doc, "对重点部位开展排查；", style="List Number")
    paragraph(doc, "建立问题台账；", style="List Number")
    paragraph(doc, "限期整改销号。", style="List Number")
    doc.add_heading("（二）监测预警", level=2)
    warning = doc.add_paragraph("完善监测预警体系，")
    warning.add_run("重点").bold = True
    warning.add_run("加强山洪灾害监测，及时发布预警信息。")
    doc.add_heading("三、保障措施", level=1)
    paragraph(doc, "工作要求", bold_lead="工作要求")
    paragraph(
        doc,
        "各单位要高度重视，加强组织领导，细化工作措施，确保各项任务落到实处。",
    )
    add_table(
        doc,
        [
            ["责任单位", "主要任务", "完成时限"],
            ["市应急管理局", "统筹协调", "2026年12月底"],
            ["各区县人民政府", "组织实施", "2026年12月底"],
        ],
    )

    doc.save(ROOT / "04-styled-plan.docx")


def text_notice():
    lines = [
        "秘密★1年",
        "×政发〔2026〕15号",
        "关于做好2026年防汛工作的通知",
        "各区县人民政府：",
        "为切实做好防汛工作，现将有关事项通知如下。",
        "一、总体要求",
        "坚持底线思维，全面排查风险隐患。",
        "（一）加强组织领导。各单位要成立领导小组。",
        "（二）压实工作责任",
        "1.排查隐患",
        "对重点部位开展拉网式排查。",
        "2.整改落实",
        "对发现问题限期整改。",
        "二、保障措施",
        "1.加强值班值守；",
        "2.强化物资储备；",
        "3.严格信息报送。",
        "各单位要于5月20日前报送工作情况。",
        "××市人民政府办公室",
        "2026年5月10日",
        "抄送：市委办公室。",
        "××市人民政府办公室 2026年5月10日印发",
    ]
    (ROOT / "05-word-export.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")


def wrapped_text():
    paragraphs = [
        "关于做好2026年防汛工作的通知",
        "各区县人民政府：",
        "为切实做好2026年防汛工作，保障人民群众生命财产安全，根据市委、市政府工作部署，结合我市实际，现将有关事项通知如下。",
        "一、总体要求",
        "坚持人民至上、生命至上，立足防大汛、抢大险、救大灾，全面排查整治各类风险隐患，确保安全度汛。",
        "（一）加强组织领导。各单位要高度重视，成立由主要负责同志任组长的防汛工作领导小组，明确责任分工，压实工作责任。",
        "（二）压实工作责任",
        "1.排查隐患",
        "对河道、水库、地质灾害隐患点开展拉网式排查，建立问题台账，实行销号管理。",
        "2.整改落实",
        "对排查发现的问题限期整改，重大隐患实行挂牌督办，确保整改到位。",
        "二、保障措施",
        "1.加强值班值守；",
        "2.强化物资储备；",
        "3.严格信息报送。",
        "各单位要于2026年5月20日前将工作情况报送市防汛抗旱指挥部办公室。",
        "××市人民政府办公室",
        "2026年5月10日",
    ]
    # 按 32 个汉字硬折行，段与段之间留空行。
    wrapped = []
    for text in paragraphs:
        line = ""
        for ch in text:
            line += ch
            if len(line) >= 32:
                wrapped.append(line)
                line = ""
        if line:
            wrapped.append(line)
        wrapped.append("")
    (ROOT / "06-wrapped.txt").write_text("\n".join(wrapped).rstrip() + "\n", encoding="utf-8")


def negative():
    doc = Document()
    paragraph(doc, "在全市防汛工作会议上的讲话")
    paragraph(doc, "同志们：")
    paragraph(doc, "今天，我们召开全市防汛工作会议，主要任务是分析形势、明确任务、压实责任。")
    paragraph(doc, "一、认清形势，切实增强做好防汛工作的责任感")
    paragraph(doc, "一、二季度全市降水偏多，防汛形势严峻。")
    paragraph(doc, "1.5亿元防汛资金已下达。")
    paragraph(doc, "2、3月份降水偏多。")
    paragraph(doc, "3.2.1 技术路线")
    paragraph(doc, "各有关单位要抓好落实。")
    paragraph(doc, "要加强值班值守")
    paragraph(doc, "二、突出重点，全力做好各项防汛工作")
    paragraph(doc, "要突出抓好水库安全度汛、山洪灾害防御、城市内涝防范等重点工作。")
    doc.save(ROOT / "07-negative.docx")


def main():
    fangxun_notice()
    approval()
    regulation()
    styled_plan()
    text_notice()
    wrapped_text()
    negative()


if __name__ == "__main__":
    main()
